from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Load the DOCX file
doc = Document('manuscript/chapter-01-assigned-to-witness.docx')

# Search for all paragraphs related to the Initiation queue passage
print('Searching for relevant paragraphs...')
para_indices = []

for i, para in enumerate(doc.paragraphs):
    if any(keyword in para.text for keyword in ['central grid flickered', 'Initiation queue', 'Three entries']):
        print(f'Paragraph {i}: {para.text}')
        para_indices.append(i)

# The new passage (as single lines per paragraph to maintain formatting)
new_passages = [
    'Then the central grid flickered.',
    '',
    'The Initiation queue populated on screen—three entries materialized like component numbers in a system ledger, stripped of color or ceremony, rendered as mere data points in Trinity\'s inventory.',
    '',
    'Elara leaned toward the screen. The names belonged to Miner Town.',
]

# Replace the old passages
if para_indices:
    print(f'\nReplacing paragraphs at indices: {para_indices}')
    
    # Get the first paragraph to update
    first_para = doc.paragraphs[para_indices[0]]
    
    # Update the first paragraph with the new content
    first_para.text = new_passages[0]
    
    # For subsequent paragraphs in para_indices, we'll update them
    for idx, para_idx in enumerate(para_indices[1:], 1):
        if idx < len(new_passages):
            doc.paragraphs[para_idx].text = new_passages[idx]
    
    print('Replacement complete')

# Save
doc.save('manuscript/chapter-01-assigned-to-witness.docx')
print('DOCX file refreshed successfully')
