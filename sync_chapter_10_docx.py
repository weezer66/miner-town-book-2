from pathlib import Path
import re

from docx import Document
from docx.shared import RGBColor


manuscript_dir = Path("manuscript")
markdown_path = manuscript_dir / "chapter-10-the-vacancy.md"
template_path = manuscript_dir / "chapter-01-assigned-to-witness.docx"
output_path = manuscript_dir / "chapter-10-the-vacancy.docx"

markdown_text = markdown_path.read_text(encoding="utf-8").strip()
blocks = markdown_text.split("\n\n")

if not blocks[0].startswith("# "):
    raise SystemExit("Chapter Markdown must begin with a level-one heading")

chapter_title = blocks[0][2:]
body_blocks = blocks[1:]

def plain_text(block: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", block.replace("\n", " "))
    if text.startswith("*") and text.endswith("*"):
        return text[1:-1]
    return text


expected_paragraphs = [chapter_title, *[plain_text(block) for block in body_blocks]]

doc = Document(template_path)
if not doc.paragraphs:
    raise SystemExit("Template must contain a title paragraph")

title_style = doc.paragraphs[0].style
body = doc._element.body

for child in list(body):
    if child.tag.endswith("}sectPr"):
        continue
    body.remove(child)

heading = doc.add_paragraph(style=title_style)
heading.add_run(chapter_title)

for block in body_blocks:
    paragraph = doc.add_paragraph()
    is_italic_block = block.strip().startswith("*") and block.strip().endswith("*")
    run = paragraph.add_run(plain_text(block))
    if is_italic_block:
        run.italic = True
        run.font.color.rgb = RGBColor(92, 107, 125)

doc.save(output_path)

actual_paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]

if actual_paragraphs != expected_paragraphs:
    raise SystemExit("Verification failed: DOCX paragraphs differ from Markdown")

print(f"Created and verified {output_path}: {len(expected_paragraphs)} paragraphs match Markdown")