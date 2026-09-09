from pathlib import Path
import re
import sys

from docx import Document


manuscript_dir = Path("manuscript")
markdown_path = manuscript_dir / "chapter-06-the-grandmaster's-briefing.md"
template_path = manuscript_dir / "chapter-01-assigned-to-witness.docx"
output_path = manuscript_dir / "chapter-06-the-grandmaster's-briefing.docx"

markdown_text = markdown_path.read_text(encoding="utf-8").strip()
blocks = markdown_text.split("\n\n")

if not blocks[0].startswith("# "):
    raise SystemExit("Chapter Markdown must begin with a level-one heading")

chapter_title = blocks[0][2:]
body_blocks = blocks[1:]
expected_paragraphs = [chapter_title, *[re.sub(r"\*\*(.*?)\*\*", r"\1", block.replace("\n", " ")) for block in body_blocks]]

if "--verify" not in sys.argv:
    doc = Document(template_path)
    title_style = doc.paragraphs[0].style
    body = doc._element.body

    for child in list(body):
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)

    heading = doc.add_paragraph(style=title_style)
    heading.add_run(chapter_title)

    for block in body_blocks:
        paragraph = doc.add_paragraph()
        paragraph.add_run(re.sub(r"\*\*(.*?)\*\*", r"\1", block.replace("\n", " ")))

    doc.save(output_path)

actual_paragraphs = [paragraph.text for paragraph in Document(output_path).paragraphs]

if actual_paragraphs != expected_paragraphs:
    raise SystemExit("Verification failed: DOCX paragraphs differ from Markdown")

action = "Verified" if "--verify" in sys.argv else "Created and verified"
print(f"{action} {output_path}: {len(expected_paragraphs)} paragraphs match Markdown")
