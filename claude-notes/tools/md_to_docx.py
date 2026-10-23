"""Convert a v2 chapter Markdown file to .docx, using the author's Chapter 1 Word file as the style template.

Usage: python md_to_docx.py CHAPTER.md TEMPLATE.docx OUT.docx

- Bold black title, a blank spacer, the italic grey epigraph and attribution, and bold subheadings.
- Inline *italic*, **bold** and ~~strikethrough~~ runs.
- A "---" scene break is kept as its own paragraph, matching the manuscript .docx files.
- A block whose lines are all upper case (ledger and record lines) becomes one paragraph per line.
- The template is opened read-only; the output is re-read and checked paragraph by paragraph.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import RGBColor

md_path, template_path, out_path = map(Path, sys.argv[1:4])

blocks = [b for b in md_path.read_text(encoding="utf-8").replace("\r", "").strip().split("\n\n") if b.strip()]
if not blocks[0].startswith("# "):
    raise SystemExit("Chapter Markdown must begin with a level-one heading")

title = blocks[0][2:].strip()
body = blocks[1:]  # "---" scene breaks stay as their own paragraph, as in the manuscript .docx files

doc = Document(template_path)
title_style = doc.paragraphs[0].style
xml_body = doc._element.body
for child in list(xml_body):
    if not child.tag.endswith("}sectPr"):
        xml_body.remove(child)

GREY = RGBColor(92, 107, 125)
BLACK = RGBColor(0, 0, 0)

p = doc.add_paragraph(style=title_style)
r = p.add_run(title)
r.bold = True
r.font.color.rgb = BLACK
doc.add_paragraph()

TOKEN = re.compile(r"(~~.+?~~|\*\*.+?\*\*|\*[^*]+?\*)")


def add_inline(paragraph, text, italic=False, color=None):
    for part in TOKEN.split(text):
        if not part:
            continue
        run_italic, run_bold, run_strike = italic, False, False
        if part.startswith("~~") and part.endswith("~~"):
            part, run_strike = part[2:-2], True
        elif part.startswith("**") and part.endswith("**"):
            part, run_bold = part[2:-2], True
        elif part.startswith("*") and part.endswith("*"):
            part, run_italic = part[1:-1], not italic
        run = paragraph.add_run(part)
        run.italic = run_italic
        run.bold = run_bold
        run.font.strike = run_strike
        if color is not None:
            run.font.color.rgb = color


def plain(t):
    return TOKEN.sub(lambda m: m.group(0).strip("*~"), t)


split_body = []
for block in body:
    lines = block.split("\n")
    if len(lines) > 1 and all(l.strip() and l.strip() == l.strip().upper() for l in lines):
        split_body.extend(lines)
    else:
        split_body.append(block)

expected = [title]
for block in split_body:
    text = block.replace("\n", " ").strip()
    para = doc.add_paragraph()
    if text.startswith("## "):
        run = para.add_run(text[3:])
        run.bold = True
        run.font.color.rgb = BLACK
        expected.append(text[3:])
    elif text.startswith("*") and text.endswith("*") and not text.startswith("**"):
        add_inline(para, text[1:-1], italic=True, color=GREY)
        expected.append(plain(text[1:-1]))
    elif text.startswith("—"):
        add_inline(para, text, italic=True, color=GREY)
        expected.append(text)
    else:
        add_inline(para, text)
        expected.append(plain(text))

doc.save(out_path)

actual = [q.text for q in Document(out_path).paragraphs if q.text.strip()]
if actual != expected:
    for i, (a, e) in enumerate(zip(actual, expected)):
        if a != e:
            raise SystemExit(f"Verification failed at paragraph {i}:\n  got:  {a[:80]}\n  want: {e[:80]}")
    raise SystemExit(f"Verification failed: {len(actual)} paragraphs vs {len(expected)} expected")
print(f"Created and verified {out_path.name}: {len(expected)} paragraphs match the Markdown")
