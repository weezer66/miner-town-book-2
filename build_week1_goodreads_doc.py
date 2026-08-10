from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUTPUT_PATH = "week 1 blogs goodreads.docx"

doc = Document()

title = doc.add_heading("Week 1 Blogs Goodreads", level=0)
title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

doc.add_paragraph("Ready-to-paste Goodreads daily posts for Week 1.")

days = [
    (
        "Day 1 - Why I Wrote Miner Town: Awakening",
        "I wanted to write a story about pressure, not perfection.\n\n"
        "In Miner Town: Awakening, the world is built so ordinary people carry impossible burdens while powerful systems stay hidden above them. The story begins in a mining settlement where survival is daily, trust is fragile, and every decision has a cost.\n\n"
        "If you enjoy character-driven speculative fiction with moral tension, that is the lane this book lives in.\n\n"
        "If you have read it, I always value honest reader reactions. Even one line about what worked for you or what did not is genuinely useful as I shape what comes next.\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 2 - A Real-World Detail Behind the Setting",
        "One thing that shaped Miner Town: Awakening was reading about how resource economies change communities over time.\n\n"
        "I kept coming back to a simple question: what happens when a whole town is organized around extraction, scarcity, and control? Not just physically, but emotionally and socially.\n\n"
        "That idea became the backbone of Miner Town. The setting is fictional, but the pressure logic is grounded: when systems reward fear and dependence, relationships become survival tools.\n\n"
        "I love stories where worldbuilding changes character choices, not just scenery. If you are the same, this book was written for that kind of reading experience.\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 3 - What Kind of Story This Is (and Is Not)",
        "Quick orientation for new readers:\n\n"
        "Miner Town: Awakening is YA dystopian with mystery and survival elements. It is not a high-magic fantasy, and it is not a pure action sprint. The heart of it is people under pressure making difficult choices.\n\n"
        "If you like:\n"
        "- moral stakes over easy heroes\n"
        "- sibling and team dynamics\n"
        "- slow-burn truth reveals\n"
        "- atmosphere with momentum\n\n"
        "you will likely feel at home in this world.\n\n"
        "If you have already read it, I am curious what you connected with most: the setting, the relationships, or the central conflict.\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 4 - One Craft Rule I Used in This Book",
        "A writing rule that helped me most in Miner Town: Awakening:\n\n"
        "Every scene must force a tradeoff.\n\n"
        "Comfort vs truth.\n"
        "Safety vs loyalty.\n"
        "Silence vs consequence.\n\n"
        "That rule kept the story from drifting and helped each chapter carry emotional weight, not just plot movement. When characters are forced to choose, readers learn who they are.\n\n"
        "I also found that tension feels stronger when each character believes they are right for different reasons. That is where most of the conflict in Miner Town comes from.\n\n"
        "If you are a writer too, what scene rule has saved you the most work in revision?\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 5 - Reader Discussion Question",
        "I would love your take on this:\n\n"
        "In a collapsing system, which matters more first, truth or trust?\n\n"
        "Miner Town: Awakening keeps returning to that question. Characters discover pieces of truth at different times, but without trust, even the right information can break a group apart.\n\n"
        "I think many dystopian stories are really stories about social trust under pressure, and that is exactly what I wanted to explore here.\n\n"
        "If you have read books in this lane recently, which ones handled this balance especially well for you? I am always adding to my reading list.\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 6 - Character Focus: Camilla",
        "Today I wanted to spotlight Camilla.\n\n"
        "She is one of my favorite characters to write because she sees patterns others miss and refuses easy comfort when truth is uncomfortable. In a world like Miner Town, that is both her strength and her burden.\n\n"
        "A lot of the emotional tension in the story comes from characters who care deeply but disagree on method. Camilla often stands at that fault line.\n\n"
        "If you enjoy stories where character intelligence is not just technical but emotional and strategic, you will probably connect with her arc.\n\n"
        "Which character type do you usually bond with first as a reader: the idealist, the strategist, the protector, or the skeptic?\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
    (
        "Day 7 - Week One Reflection + Small Excerpt",
        "Thank you for the thoughtful responses this week. I wanted to close with a short line from Miner Town: Awakening that reflects the core mood:\n\n"
        "Some worlds run on the bones of the one before. This one was beginning to choke on them.\n\n"
        "This book was built around pressure, loyalty, and the cost of waking up to truth when truth does not come with safety.\n\n"
        "Next week I will share a short note on how I revised tension without adding extra plot noise, since several readers asked about process.\n\n"
        "If you are reading now or planning to, I appreciate honest feedback more than polished feedback. Real reactions help me write better books.\n\n"
        "Disclosure: I am the author of Miner Town: Awakening."
    ),
]

for heading, body in days:
    doc.add_heading(heading, level=1)
    for para in body.split("\n\n"):
        doc.add_paragraph(para)

# Minimal formatting consistency
for paragraph in doc.paragraphs:
    if paragraph.style.name == "Normal":
        paragraph.paragraph_format.space_after = 8


doc.save(OUTPUT_PATH)
print(f"Wrote {OUTPUT_PATH}")
