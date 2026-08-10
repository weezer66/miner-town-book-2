#!/usr/bin/env python3
"""
Send launch-day emails in bursts and maintain a side tracker workbook.

Behavior:
- Sends launch-day emails using launch_day_campaign.py in rounds.
- Each round sends up to --batch recipients.
- Sleeps --cooldown-minutes between rounds.
- Uses safe per-email pacing and daily caps by default.
- Rebuilds tracker workbook each round with per-email status:
  SENT / PENDING / SUPPRESSED.
"""

import argparse
import csv
import os
import re
import subprocess
import sys
import time
from datetime import datetime

from openpyxl import Workbook, load_workbook


HERE = os.path.dirname(os.path.abspath(__file__))
SHEET_PATH = os.path.join(HERE, "followups.xlsx")
SUPPRESSION_PATH = os.path.join(HERE, "suppression.txt")
LAUNCH_LOG_PATH = os.path.join(HERE, "launch_day_sent_log.txt")
TRACKER_PATH = os.path.join(HERE, "launch_day_send_tracker.xlsx")


def load_suppressed(path):
    out = set()
    if not os.path.exists(path):
        return out
    with open(path, encoding="utf-8", errors="ignore") as f:
        for line in f:
            t = line.strip().lower()
            if not t or t.startswith("#"):
                continue
            out.add(t)
    return out


def load_followup_rows(path):
    if not os.path.exists(path):
        raise FileNotFoundError(f"Sheet not found: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Follow-ups"] if "Follow-ups" in wb.sheetnames else wb.active

    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=c).value
        if v:
            headers[str(v).strip()] = c

    if "email" not in headers:
        wb.close()
        raise ValueError("followups.xlsx is missing 'email' column")

    rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        email_raw = r[headers["email"] - 1] if headers.get("email") else ""
        email = (str(email_raw).strip().lower() if email_raw else "")
        if not email or "@" not in email:
            continue
        responded = ""
        never = ""
        if headers.get("responded"):
            v = r[headers["responded"] - 1]
            responded = (str(v).strip().lower() if v is not None else "")
        if headers.get("Never Follow-up?"):
            v = r[headers["Never Follow-up?"] - 1]
            never = (str(v).strip().lower() if v is not None else "")

        # Keep first row for duplicate emails and ignore repeats.
        rows.setdefault(
            email,
            {
                "email": email,
                "responded": responded,
                "never_followup": never,
            },
        )

    wb.close()
    return rows


def load_campaign_sent_map(log_path, campaign_id):
    sent = {}
    if not os.path.exists(log_path):
        return sent

    # Backward-compatible parser: supports legacy files where separators were
    # written as literal "\\n" text instead of newline characters.
    with open(log_path, encoding="utf-8", errors="ignore") as f:
        raw = f.read().replace("\\n", "\n")

    reader = csv.reader(raw.splitlines())
    for row in reader:
        if not row or row[0] == "campaign_id":
            continue
        if len(row) < 3:
            continue
        cid = row[0].strip()
        email = row[1].strip().lower()
        sent_at = row[2].strip()
        if cid != campaign_id:
            continue
        if email:
            sent[email] = sent_at
    return sent


def write_tracker(path, rows_map, suppressed, sent_map, campaign_id):
    wb = Workbook()
    ws = wb.active
    ws.title = "Launch Day Status"
    ws.append(
        [
            "email",
            "responded",
            "never_followup",
            "suppressed",
            "status",
            "sent_at",
            "campaign_id",
            "updated_at",
        ]
    )

    stamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for email in sorted(rows_map.keys()):
        row = rows_map[email]
        is_suppressed = email in suppressed
        sent_at = sent_map.get(email, "")
        if is_suppressed:
            status = "SUPPRESSED"
        elif sent_at:
            status = "SENT"
        else:
            status = "PENDING"

        ws.append(
            [
                email,
                row["responded"],
                row["never_followup"],
                "yes" if is_suppressed else "no",
                status,
                sent_at,
                campaign_id,
                stamp,
            ]
        )

    wb.save(path)


def compute_counts(rows_map, suppressed, sent_map):
    total = len(rows_map)
    suppressed_count = 0
    sent_count = 0
    pending_count = 0

    for email in rows_map.keys():
        if email in suppressed:
            suppressed_count += 1
        elif email in sent_map:
            sent_count += 1
        else:
            pending_count += 1

    return {
        "total": total,
        "suppressed": suppressed_count,
        "sent": sent_count,
        "pending": pending_count,
    }


def parse_sent_now(output_text):
    m = re.search(r"Sent now:\s*(\d+)", output_text or "")
    return int(m.group(1)) if m else 0


def run_launch_round(
    pyexe,
    campaign_id,
    batch_size,
    daily_limit,
    subject,
    sleep_min,
    sleep_max,
    override_guardrail,
):
    cmd = [
        pyexe,
        os.path.join(HERE, "launch_day_campaign.py"),
        "--send",
        "--campaign-id",
        campaign_id,
        "--include-responded",
        "all",
        "--include-never",
        "--standalone",
        "--limit",
        str(batch_size),
        "--daily-limit",
        str(daily_limit),
        "--sleep-min",
        str(sleep_min),
        "--sleep-max",
        str(sleep_max),
        "--subject",
        subject,
    ]
    if override_guardrail:
        cmd.append("--override-guardrail")
    proc = subprocess.run(cmd, text=True, capture_output=True, cwd=HERE)
    if proc.stdout:
        print(proc.stdout, end="")
    if proc.stderr:
        print(proc.stderr, end="", file=sys.stderr)
    return proc.returncode, parse_sent_now(proc.stdout)


def main():
    ap = argparse.ArgumentParser(
        description="Send launch-day emails in bursts and update a side tracker workbook."
    )
    ap.add_argument("--campaign-id", default=f"launch-all-{datetime.now():%Y-%m-%d}")
    ap.add_argument("--batch", type=int, default=2, help="Emails per burst round")
    ap.add_argument(
        "--cooldown-minutes",
        type=int,
        default=20,
        help="Minutes to wait between rounds",
    )
    ap.add_argument(
        "--daily-limit",
        type=int,
        default=100,
        help="Daily cap for this campaign",
    )
    ap.add_argument(
        "--sleep-min",
        type=float,
        default=20.0,
        help="Min seconds between individual sends inside each round.",
    )
    ap.add_argument(
        "--sleep-max",
        type=float,
        default=45.0,
        help="Max seconds between individual sends inside each round.",
    )
    ap.add_argument(
        "--override-guardrail",
        action="store_true",
        help="Pass --override-guardrail to launch_day_campaign.py (not recommended).",
    )
    ap.add_argument(
        "--tracker",
        default=TRACKER_PATH,
        help="Path to tracker xlsx",
    )
    ap.add_argument(
        "--python",
        default=sys.executable,
        help="Python executable to invoke launch_day_campaign.py",
    )
    ap.add_argument(
        "--subject",
        default="My new novel, Miner Town: Awakening, is out now",
        help="Standalone launch subject",
    )
    args = ap.parse_args()

    if args.batch <= 0:
        sys.exit("--batch must be > 0")
    if args.cooldown_minutes < 0:
        sys.exit("--cooldown-minutes must be >= 0")
    if args.daily_limit <= 0:
        sys.exit("--daily-limit must be > 0")
    if args.sleep_min < 0 or args.sleep_max < 0 or args.sleep_max < args.sleep_min:
        sys.exit("Sleep bounds must be non-negative and sleep-max >= sleep-min")
    if (
        not args.override_guardrail
        and (args.sleep_min < 15.0 or args.sleep_max < 15.0)
    ):
        sys.exit(
            "Guardrail blocked burst run: --sleep-min and --sleep-max must both be >= 15 seconds. "
            "Use slower pacing or pass --override-guardrail."
        )

    if not os.environ.get("SMTP_USER") or not os.environ.get("SMTP_PASS"):
        sys.exit("Set SMTP_USER and SMTP_PASS first.")

    rows_map = load_followup_rows(SHEET_PATH)
    suppressed = load_suppressed(SUPPRESSION_PATH)
    sent_map = load_campaign_sent_map(LAUNCH_LOG_PATH, args.campaign_id)
    write_tracker(args.tracker, rows_map, suppressed, sent_map, args.campaign_id)
    counts = compute_counts(rows_map, suppressed, sent_map)

    print(f"Campaign: {args.campaign_id}")
    print(f"Tracker:  {args.tracker}")
    print(
        f"Initial totals -> total={counts['total']} sent={counts['sent']} "
        f"pending={counts['pending']} suppressed={counts['suppressed']}"
    )

    round_no = 0
    while True:
        sent_map = load_campaign_sent_map(LAUNCH_LOG_PATH, args.campaign_id)
        counts = compute_counts(rows_map, suppressed, sent_map)
        write_tracker(args.tracker, rows_map, suppressed, sent_map, args.campaign_id)

        if counts["pending"] <= 0:
            print("All non-suppressed recipients are marked SENT for this campaign.")
            break

        round_no += 1
        print(
            f"\n=== ROUND {round_no} === pending={counts['pending']} "
            f"(batch={args.batch})"
        )
        rc, sent_now = run_launch_round(
            pyexe=args.python,
            campaign_id=args.campaign_id,
            batch_size=args.batch,
            daily_limit=args.daily_limit,
            subject=args.subject,
            sleep_min=args.sleep_min,
            sleep_max=args.sleep_max,
            override_guardrail=args.override_guardrail,
        )

        sent_map = load_campaign_sent_map(LAUNCH_LOG_PATH, args.campaign_id)
        counts = compute_counts(rows_map, suppressed, sent_map)
        write_tracker(args.tracker, rows_map, suppressed, sent_map, args.campaign_id)
        print(
            f"ROUND {round_no} result -> sent_now={sent_now}, exit={rc}, "
            f"remaining_pending={counts['pending']}"
        )

        if rc != 0:
            print("Stopping due to non-zero exit from launch_day_campaign.py")
            break
        if sent_now <= 0:
            print("Stopping because round sent 0 emails (nothing sendable or blocked).")
            break
        if counts["pending"] <= 0:
            print("Campaign complete.")
            break

        if args.cooldown_minutes > 0:
            print(f"Cooldown: sleeping {args.cooldown_minutes} minute(s)...")
            time.sleep(args.cooldown_minutes * 60)

    # Final refresh of tracker for certainty.
    sent_map = load_campaign_sent_map(LAUNCH_LOG_PATH, args.campaign_id)
    write_tracker(args.tracker, rows_map, suppressed, sent_map, args.campaign_id)
    final = compute_counts(rows_map, suppressed, sent_map)
    print(
        f"\nFINAL -> total={final['total']} sent={final['sent']} "
        f"pending={final['pending']} suppressed={final['suppressed']}"
    )


if __name__ == "__main__":
    main()
