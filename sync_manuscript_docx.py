from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    current_lines: list[str] = []

    def flush_paragraph() -> None:
        if not current_lines:
            return
        content = " ".join(part.strip() for part in current_lines if part.strip())
        if content:
            doc.add_paragraph(content)
        current_lines.clear()

    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            doc.add_heading(stripped[2:].strip(), level=1)
            continue

        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_heading(stripped[3:].strip(), level=2)
            continue

        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:].strip(), level=3)
            continue

        if stripped.startswith("- ") or stripped.startswith("* "):
            flush_paragraph()
            doc.add_paragraph(stripped[2:].strip())
            continue

        current_lines.append(line.strip())

    flush_paragraph()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"Synced {md_path} -> {docx_path}")


def sync_all_manuscript_files(manuscript_dir: Path) -> None:
    for md_path in sorted(manuscript_dir.glob("*.md")):
        docx_path = md_path.with_suffix(".docx")
        markdown_to_docx(md_path, docx_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Sync manuscript Markdown chapters to DOCX files.")
    parser.add_argument("--md", type=Path, help="Optional single Markdown file to sync.")
    parser.add_argument("--docx", type=Path, help="Optional DOCX output path for a single file sync.")
    parser.add_argument("--dir", type=Path, default=Path("manuscript"), help="Directory containing manuscript markdown files.")
    args = parser.parse_args()

    if args.md and args.docx:
        markdown_to_docx(args.md, args.docx)
        return

    sync_all_manuscript_files(args.dir)


if __name__ == "__main__":
    main()
